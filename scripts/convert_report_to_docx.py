#!/usr/bin/env python3
"""
Convert DIPLOMA_THESIS_REPORT.md to academic standard DOCX format (ДСТУ / МОН України).
- Font: Times New Roman, 14pt (tables/code 10-12pt)
- Line spacing: 1.5
- Margins: Left 30mm, Right 10mm, Top 20mm, Bottom 20mm
- Paragraph indent: 1.25 cm
- Alignment: Justified
- Headings: Bold, Page breaks before major sections
"""

import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips (1/20 of a point)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    """Set background color for a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Set subtle border to entire table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_inline_markdown(p, text, base_font_size=Pt(14), italic_override=False, code_mode=False):
    """Parse bold, italic, and inline code formatting."""
    # Pattern to match bold, italic, code, and plain text tokens
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = base_font_size
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = base_font_size
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(40, 40, 140)
        else:
            run = p.add_run(token)
            run.font.name = 'Times New Roman' if not code_mode else 'Consolas'
            run.font.size = base_font_size
            if italic_override:
                run.italic = True
            run.font.color.rgb = RGBColor(0, 0, 0)

def main():
    md_path = Path("docs/DIPLOMA_THESIS_REPORT.md")
    docx_path = Path("docs/DIPLOMA_THESIS_REPORT.docx")
    
    if not md_path.exists():
        print(f"Error: {md_path} not found")
        return 1

    doc = docx.Document()

    # Configure page setup (A4, Academic Margins in Ukraine)
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(10)

    # Base style configurations
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(lines):
        if not lines:
            return
        parsed_rows = []
        for line in lines:
            if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', line):
                continue  # separator row
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if cells:
                parsed_rows.append(cells)
        
        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="999999", sz="4")

        for row_idx, row_data in enumerate(parsed_rows):
            is_header = (row_idx == 0)
            for col_idx in range(num_cols):
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = table.cell(row_idx, col_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                if is_header:
                    set_cell_background(cell, "EAEAEA")

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.first_line_indent = Pt(0)
                
                format_inline_markdown(p, cell_text, base_font_size=Pt(11))
                if is_header:
                    for r in p.runs:
                        r.bold = True

        # Add spacing after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)
        sp.paragraph_format.line_spacing = 1.0

    def flush_code(lines):
        if not lines:
            return
        # Create a single cell table for code block with subtle background
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="CCCCCC", sz="4")
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Pt(0)

        for i, cl in enumerate(lines):
            run = p.add_run(cl + ('\n' if i < len(lines) - 1 else ''))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(30, 30, 30)

        # Spacing after code block
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)
        sp.paragraph_format.line_spacing = 1.0

    i = 0
    is_title_page = True

    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')

        # Code block toggle
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table detection
        if '|' in line and line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Page separator in markdown
        if line.strip() == '---':
            doc.add_page_break()
            is_title_page = False
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_inline_markdown(p, text, base_font_size=Pt(16))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        if line.startswith('## '):
            text = line[3:].strip()
            # If it's a major section like РОЗДІЛ, ВСТУП, ВИСНОВКИ, АНОТАЦІЯ, add page break
            is_major_section = any(k in text.upper() for k in ['РОЗДІЛ', 'ВСТУП', 'ВИСНОВКИ', 'СПИСОК', 'АНОТАЦІЯ', 'ANNOTATION', 'ЗМІСТ', 'ДОДАТОК', 'ЧЕРНІВЕЦЬКИЙ'])
            if is_major_section and not is_title_page:
                doc.add_page_break()

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_major_section else WD_ALIGN_PARAGRAPH.LEFT
            format_inline_markdown(p, text, base_font_size=Pt(14))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Pt(35.4)  # 1.25 cm
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_inline_markdown(p, text, base_font_size=Pt(14))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Pt(35.4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_inline_markdown(p, text, base_font_size=Pt(14))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # Bullet list item
        if line.strip().startswith(('- ', '* ')):
            text = line.strip()[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(35.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run("• ")
            run.bold = True
            format_inline_markdown(p, text, base_font_size=Pt(14))
            i += 1
            continue

        # Numbered list item
        num_match = re.match(r'^\s*(\d+[\.\)])\s+(.*)$', line)
        if num_match:
            prefix = num_match.group(1)
            text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(35.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(prefix + " ")
            run.bold = True
            format_inline_markdown(p, text, base_font_size=Pt(14))
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5

        # Check if title page lines
        if is_title_page:
            p.paragraph_format.first_line_indent = Pt(0)
            if "МІНІСТЕРСТВО" in line or "ЧЕРНІВЕЦЬКИЙ" in line or "Навчально-науковий" in line or "Кафедра" in line or "Чернівці" in line or "СИСТЕМА" in line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif "Виконав" in line or "студент" in line or "Горинюк" in line or "Керівник" in line or "кандидат" in line or "Дворжак" in line:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.paragraph_format.first_line_indent = Pt(35.4)  # 1.25 cm indent
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        format_inline_markdown(p, line, base_font_size=Pt(14))
        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    doc.save(str(docx_path))
    print(f"Successfully generated: {docx_path}")
    return 0

if __name__ == "__main__":
    exit(main())
